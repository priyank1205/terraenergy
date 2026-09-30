"""
Energy Institute — Statistical Review of World Energy 2026 (data through 2025).

Parses the "all data" workbook: country time series in physical units, regional and
world totals, and the 2025 bilateral LNG / pipeline-gas trade matrices.
"""

from __future__ import annotations

from dataclasses import dataclass, field

import openpyxl

from pipeline.lib.countries import EI_REGION_ROWS, iso3_for_ei

# metric key -> (sheet name, unit)
SERIES_SHEETS = {
    "oil_prod_kbd": ("Oil Production - barrels", "kb/d"),
    "oil_cons_kbd": ("Oil Consumption - barrels", "kb/d"),
    "gas_prod_bcm": ("Gas Production - Bcm", "bcm"),
    "gas_cons_bcm": ("Gas Consumption - Bcm", "bcm"),
    "coal_prod_mt": ("Coal Production - mt", "Mt"),
    "coal_prod_ej": ("Coal Production - EJ", "EJ"),
    "coal_cons_ej": ("Coal Consumption - EJ", "EJ"),
    "oil_cons_ej": ("Oil Consumption - EJ", "EJ"),
    "gas_cons_ej": ("Gas Consumption - EJ", "EJ"),
    "nuclear_ej": ("Nuclear Consumption - EJ", "EJ"),
    "hydro_ej": ("Hydro Consumption - EJ", "EJ"),
    "renew_ej": ("Renewables Consumption -EJ", "EJ"),
    "tes_ej": ("Total Energy Supply (TES) -EJ", "EJ"),
    "tes_pc_gj": ("TES per Capita", "GJ/person"),
    "co2_mt": ("CO2 from Energy", "Mt CO2"),
    "elec_twh": ("Electricity Generation - TWh", "TWh"),
    "nuclear_twh": ("Nuclear Generation - TWh", "TWh"),
    "hydro_twh": ("Hydro Generation - TWh", "TWh"),
    "solar_twh": ("Solar Generation - TWh", "TWh"),
    "wind_twh": ("Wind Generation - TWh", "TWh"),
    "other_ren_twh": ("Geo Biomass Other - TWh", "TWh"),
    "refining_cap_kbd": ("Oil refinery - capacity", "kb/d"),
    "refining_thru_kbd": ("Oil refinery - throughput", "kb/d"),
    "solar_cap_mw": ("Solar Installed Capacity", "MW"),
    "wind_cap_mw": ("Wind Installed Capacity", "MW"),
    "lng_imp_bcm": ("Gas - LNG imports bcm", "bcm"),
    "lng_exp_bcm": ("Gas - LNG exports bcm", "bcm"),
}

FIRST_YEAR = 2000


@dataclass
class EIData:
    years: list[int]
    # metric -> iso3 -> {year: value}
    countries: dict[str, dict[str, dict[int, float]]] = field(default_factory=dict)
    # metric -> region code ("WORLD", "NAM", …) -> {year: value}
    regions: dict[str, dict[str, dict[int, float]]] = field(default_factory=dict)
    # metric -> {"OECD": {...}, "OPEC": {...}, "EU": {...}}
    groups: dict[str, dict[str, dict[int, float]]] = field(default_factory=dict)
    # aggregate "Other …" rows, kept so gaps are visible: metric -> label -> {year: v}
    others: dict[str, dict[str, dict[int, float]]] = field(default_factory=dict)
    lng_matrix: dict = field(default_factory=dict)
    pipe_matrix: dict = field(default_factory=dict)
    oil_trade: dict = field(default_factory=dict)
    prices: dict = field(default_factory=dict)


def _num(v):
    if isinstance(v, (int, float)):
        return float(v)
    return None


def _header(rows):
    for i, r in enumerate(rows[:8]):
        if sum(isinstance(x, int) and 1900 < x < 2100 for x in r[1:]) >= 3:
            return i, r
    raise ValueError("no year header")


def _parse_series(ws) -> tuple[dict, dict, dict, dict]:
    rows = list(ws.iter_rows(values_only=True))
    hi, hdr = _header(rows)
    # The year header repeats the latest year for growth/share columns; keep the first.
    col_for_year: dict[int, int] = {}
    for idx, y in enumerate(hdr):
        if isinstance(y, int) and FIRST_YEAR <= y <= 2100 and y not in col_for_year:
            col_for_year[y] = idx
    countries, regions, groups, others = {}, {}, {}, {}
    for r in rows[hi + 1:]:
        label = r[0]
        if not isinstance(label, str):
            continue
        name = label.strip()
        vals = {y: _num(r[c]) for y, c in col_for_year.items() if c < len(r)}
        vals = {y: v for y, v in vals.items() if v is not None}
        if not vals:
            continue
        if name in EI_REGION_ROWS:
            regions[EI_REGION_ROWS[name]] = vals
        elif name.startswith("of which: OECD"):
            groups["OECD"] = vals
        elif name == "Non-OECD":
            groups["NON_OECD"] = vals
        elif name == "OPEC":
            groups["OPEC"] = vals
        elif name.startswith("European Union"):
            groups["EU"] = vals
        else:
            iso = iso3_for_ei(name)
            if iso:
                countries[iso] = vals
            elif name.startswith(("Other", "Central America", "Eastern Africa", "Middle Africa",
                                  "Western Africa", "Rest of")):
                others[name] = vals
    return countries, regions, groups, others


# Regional subtotal rows inside the trade matrices (not importers in their own right).
_MATRIX_SUBTOTALS = {
    "north america", "s. & cent. america", "europe", "middle east & africa", "asia pacific",
    "total exports", "cis", "middle east", "africa", "total imports",
}


def _parse_matrix(ws) -> dict:
    """EI trade matrix: rows = importers ('To'), columns = exporters ('From')."""
    rows = list(ws.iter_rows(values_only=True))
    header_idx = next(i for i, r in enumerate(rows) if isinstance(r[0], str) and r[0].strip().lower().startswith("to"))
    top = rows[header_idx - 1] if header_idx > 0 else [None] * len(rows[header_idx])
    exporters = []
    for c in range(1, len(rows[header_idx])):
        parts = [p for p in (top[c] if c < len(top) else None, rows[header_idx][c]) if isinstance(p, str)]
        label = " ".join(" ".join(parts).split()) if parts else None
        if label and label.startswith("From "):
            label = label[5:]
        exporters.append(label)
    cells = []
    for r in rows[header_idx + 1:]:
        if not isinstance(r[0], str):
            continue
        importer = " ".join(r[0].split())
        if importer.lower().startswith(("source", "note", "*", "^", " *")):
            break
        if importer.lower() in _MATRIX_SUBTOTALS:
            continue
        for c, exp in enumerate(exporters, start=1):
            if not exp or exp.lower().startswith("total"):
                continue
            v = _num(r[c]) if c < len(r) else None
            if v and v > 0:
                cells.append({"from": exp, "to": importer, "bcm": v})
    return {"cells": cells}


def _parse_oil_trade(ws) -> dict:
    """'Oil - Trade movements in 24-25': crude/product imports & exports by region, 2024–25."""
    rows = list(ws.iter_rows(values_only=True))
    out = {"mt": {}, "kbd": {}}
    block = "mt"
    for r in rows[4:]:
        if isinstance(r[0], str) and "barrels" in r[0].lower():
            block = "kbd"
            continue
        if not isinstance(r[0], str) or not isinstance(r[1], (int, float)):
            continue
        name = r[0].strip()
        out[block][name] = {
            2024: {"crude_imp": r[1], "prod_imp": r[2], "crude_exp": r[3], "prod_exp": r[4]},
            2025: {"crude_imp": r[5], "prod_imp": r[6], "crude_exp": r[7], "prod_exp": r[8]},
        }
    return out


def _parse_prices(wb) -> dict:
    prices: dict[str, dict[int, float]] = {}
    ws = wb["Spot crude prices"]
    for r in ws.iter_rows(min_row=5, values_only=True):
        if isinstance(r[0], int) and r[0] >= FIRST_YEAR:
            for key, col in (("dubai", 1), ("brent", 2), ("wti", 4)):
                v = _num(r[col])
                if v is not None:
                    prices.setdefault(key, {})[r[0]] = v
    ws = wb["Gas Prices "]
    for r in ws.iter_rows(min_row=5, values_only=True):
        if isinstance(r[0], int) and r[0] >= FIRST_YEAR:
            for key, col in (("lng_japan", 1), ("jkm", 3), ("nbp", 8), ("ttf", 9), ("henry_hub", 10)):
                v = _num(r[col]) if col < len(r) else None
                if v is not None:
                    prices.setdefault(key, {})[r[0]] = v
    ws = wb["Coal & Uranium - Prices"]
    rows = list(ws.iter_rows(values_only=True))
    hdr = rows[3]
    cols = {}
    for i, x in enumerate(hdr):
        if isinstance(x, str) and "Australia" in x:
            cols["coal_australia"] = i
        elif isinstance(x, str) and "Northwest Europe" in x:
            cols["coal_nw_europe"] = i
    for r in rows[4:]:
        if isinstance(r[0], int) and r[0] >= FIRST_YEAR:
            for key, col in cols.items():
                v = _num(r[col])
                if v is not None:
                    prices.setdefault(key, {})[r[0]] = v
    return prices


def load(path) -> EIData:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    years = list(range(FIRST_YEAR, 2026))
    data = EIData(years=years)
    for metric, (sheet, _unit) in SERIES_SHEETS.items():
        c, r, g, o = _parse_series(wb[sheet])
        data.countries[metric] = c
        data.regions[metric] = r
        data.groups[metric] = g
        data.others[metric] = o
    data.lng_matrix = _parse_matrix(wb["Gas trade 2025 - LNG"])
    data.pipe_matrix = _parse_matrix(wb["Gas trade 2025 - pipeline"])
    data.oil_trade = _parse_oil_trade(wb["Oil - Trade movements in 24-25"])
    data.prices = _parse_prices(wb)
    return data
